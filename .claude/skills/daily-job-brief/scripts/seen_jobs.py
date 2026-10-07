"""Seen-jobs log for the daily job brief. Keeps every job ever shown so no job repeats.

Usage:
  python seen_jobs.py check <candidates.json> [--log PATH]
      Prints JSON: {"new": [...], "duplicates": [{"job": ..., "reason": ...}]}
  python seen_jobs.py add <jobs.json> [--log PATH]
      Appends jobs to the log (skips any already present).
  python seen_jobs.py seed <file.xlsx> [<file.xlsx> ...] [--log PATH]
      Imports Company / Role / Job link rows from existing tracker or search files.

Each job is a dict with: company, role, city, job_url, and optionally match_pct, decision.
Default log: projects/india-job-search/seen-jobs.csv (relative to the workspace root).
"""

import csv
import json
import re
import sys
from datetime import date
from pathlib import Path
from urllib.parse import parse_qsl, urlsplit

DEFAULT_LOG = Path("projects/india-job-search/seen-jobs.csv")
FIELDS = ["first_seen", "company", "role", "city", "job_url", "url_key", "job_key", "match_pct", "decision"]

COMPANY_NOISE = {"pvt", "private", "ltd", "limited", "llp", "inc", "corp", "corporation", "co", "india", "the"}
ROLE_NOISE = {"remote", "hybrid", "onsite", "on", "site", "wfh", "work", "from", "home", "fulltime", "full", "time"}
# Query parameters that identify a specific job posting; all others are tracking noise.
ID_PARAMS = re.compile(r"(job|id|jk|req|gh_jid|posting)", re.I)


def words(text):
    text = re.sub(r"\(.*?\)|\[.*?\]", " ", str(text or "").lower())
    return re.sub(r"[^a-z0-9]+", " ", text).split()


def url_key(url):
    url = str(url or "").strip()
    if not url or url in {"-", "—"}:
        return ""
    parts = urlsplit(url if "://" in url else "https://" + url)
    host = parts.netloc.lower().removeprefix("www.")
    path = parts.path.rstrip("/").lower()
    params = sorted((k.lower(), v) for k, v in parse_qsl(parts.query) if ID_PARAMS.search(k))
    query = "&".join(f"{k}={v}" for k, v in params)
    return f"{host}{path}" + (f"?{query}" if query else "")


def job_key(company, role):
    c = " ".join(w for w in words(company) if w not in COMPANY_NOISE)
    r = " ".join(w for w in words(role) if w not in ROLE_NOISE)
    return f"{c}|{r}" if c and r else ""


def load_log(log):
    if not log.exists():
        return []
    with log.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def split_new(jobs, log):
    seen_urls, seen_keys = set(), set()
    for row in load_log(log):
        if row["url_key"]:
            seen_urls.add(row["url_key"])
        if row["job_key"]:
            seen_keys.add(row["job_key"])
    new, dups = [], []
    for job in jobs:
        u, k = url_key(job.get("job_url")), job_key(job.get("company"), job.get("role"))
        if u and u in seen_urls:
            dups.append({"job": job, "reason": "same job link seen before"})
        elif k and k in seen_keys:
            dups.append({"job": job, "reason": "same company and role seen before"})
        else:
            new.append(job)
        # Also catches duplicates inside the same batch (one role posted in several cities).
        if u:
            seen_urls.add(u)
        if k:
            seen_keys.add(k)
    return new, dups


def add(jobs, log):
    new, _ = split_new(jobs, log)
    log.parent.mkdir(parents=True, exist_ok=True)
    write_header = not log.exists()
    with log.open("a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        if write_header:
            writer.writeheader()
        for job in new:
            writer.writerow({
                "first_seen": job.get("first_seen") or date.today().isoformat(),
                "company": job.get("company", ""),
                "role": job.get("role", ""),
                "city": job.get("city", ""),
                "job_url": job.get("job_url", ""),
                "url_key": url_key(job.get("job_url")),
                "job_key": job_key(job.get("company"), job.get("role")),
                "match_pct": job.get("match_pct", ""),
                "decision": job.get("decision", ""),
            })
    return len(new)


def read_xlsx(path):
    from openpyxl import load_workbook

    jobs = []
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows, [])]
        if "company" not in header or "role" not in header:
            continue
        col = {name: header.index(name) for name in header if name}
        for row in rows:
            get = lambda name: row[col[name]] if name in col and col[name] < len(row) else ""
            if not get("company"):
                continue
            jobs.append({
                "company": get("company"),
                "role": get("role"),
                "city": get("city / work mode"),
                "job_url": get("job link"),
                "match_pct": get("jd match %"),
                "decision": get("status") or "seeded",
                "first_seen": str(get("date added") or "")[:10] or None,
            })
    return jobs


def main(argv):
    log = DEFAULT_LOG
    if "--log" in argv:
        i = argv.index("--log")
        log = Path(argv[i + 1])
        del argv[i:i + 2]
    if len(argv) < 2:
        sys.exit(__doc__)
    cmd, files = argv[0], argv[1:]
    if cmd == "seed":
        jobs = [job for path in files for job in read_xlsx(path)]
        print(f"seeded {add(jobs, log)} of {len(jobs)} rows into {log}")
        return
    jobs = json.loads(Path(files[0]).read_text(encoding="utf-8"))
    if cmd == "check":
        new, dups = split_new(jobs, log)
        print(json.dumps({"new": new, "duplicates": dups}, indent=2, ensure_ascii=False, default=str))
    elif cmd == "add":
        print(f"added {add(jobs, log)} of {len(jobs)} jobs to {log}")
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
